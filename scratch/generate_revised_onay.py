# -*- coding: utf-8 -*-
"""
Akdeniz Üniversitesi İnşaat Mühendisliği Bölümü
Prof. Dr. Ramazan ÖZÇELİK ve Prof. Dr. Ferhat ERDAL Onay/Revizyon Programı Üreticisi
"""

import sys, os, copy, openpyxl
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0, os.path.abspath('.'))

# Mevcut ana motor fonksiyonlarını içe aktar
from generate_schedule import (
    load_and_schedule, TEACHERS, TIME_SLOTS, DAYS, COLOR_SCHEME,
    SECOND_EDUCATION_COURSES, populate_teacher_sheet, create_personal_excel,
    create_individual_teacher_excels, to_ascii
)

def generate_revised_onay_schedules():
    output_dir = "OnayIcinDersProgramlari"
    os.makedirs(output_dir, exist_ok=True)
    
    # 1. Mevcut programı yükle
    teacher_schedule, all_students_info = load_and_schedule()
    
    # 2. PROF. DR. FERHAT ERDAL (Şube 9) REVİZYONU
    # Amaç: Cuma = 0 saat, Pazartesi = Sadece İNM 403 (1 saat), 
    # Salı = 8 saat, Çarşamba = 8 saat, Perşembe = 6 saat.
    # Önce Ferhat Hoca'nın Pazartesi ve Cuma lisansüstü derslerini temizle
    room_9 = TEACHERS[9]['room']
    
    # Mevcut yerleşimde Ferhat Hoca'nın derslerini sıfırlayıp yeniden kurgulayalım:
    # Lisans ve seçmeli dersleri koru:
    lisans_secmeli_9 = []
    for d in range(5):
        for sl in range(9):
            e = teacher_schedule[9][d][sl]
            if e and e['type'] in ('Lisans', 'Seçmeli'):
                lisans_secmeli_9.append((d, sl, e))
            teacher_schedule[9][d][sl] = None
            
    for d, sl, e in lisans_secmeli_9:
        teacher_schedule[9][d][sl] = e

    # Uzmanlık Alan Dersi: FBE 9901 (8 saat = 2 x 4 saat)
    # 1. Blok: Salı 08:30 - 12:20 (Slot 0, 1, 2, 3)
    for sl in [0, 1, 2, 3]:
        teacher_schedule[9][1][sl] = {
            'sube': '9', 'code': 'FBE 9901', 'name': 'Uzmanlık Alan Dersi',
            'day': 1, 'slot': sl, 'start': TIME_SLOTS[sl][0], 'end': TIME_SLOTS[sl][1],
            'room': room_9, 'uyg': 0, 'ortak': 0, 'ckd': 0, 'prog': '16',
            'type': 'Uzmanlık', 'student_name': 'OSMAN CAN KAYA'
        }
    # 2. Blok: Çarşamba 13:30 - 17:20 (Slot 5, 6, 7, 8)
    for sl in [5, 6, 7, 8]:
        teacher_schedule[9][2][sl] = {
            'sube': '9', 'code': 'FBE 9901', 'name': 'Uzmanlık Alan Dersi',
            'day': 2, 'slot': sl, 'start': TIME_SLOTS[sl][0], 'end': TIME_SLOTS[sl][1],
            'room': room_9, 'uyg': 0, 'ortak': 0, 'ckd': 0, 'prog': '16',
            'type': 'Uzmanlık', 'student_name': 'OSMAN CAN KAYA'
        }

    # Danışmanlıklar:
    # Danışmanlık 1: Osman Can Kaya -> Salı 13:30 - 14:20 (Slot 5)
    teacher_schedule[9][1][5] = {
        'sube': '9', 'code': 'FBE 9903', 'name': 'Danışmanlık',
        'day': 1, 'slot': 5, 'start': TIME_SLOTS[5][0], 'end': TIME_SLOTS[5][1],
        'room': room_9, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '16',
        'type': 'Danışmanlık', 'student_name': 'OSMAN CAN KAYA'
    }
    # Danışmanlık 2: Adriana Lizzette Cruz Quintanilla -> Salı 14:30 - 15:20 (Slot 6)
    teacher_schedule[9][1][6] = {
        'sube': '109', 'code': 'FBE 9913', 'name': 'Danışmanlık',
        'day': 1, 'slot': 6, 'start': TIME_SLOTS[6][0], 'end': TIME_SLOTS[6][1],
        'room': room_9, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '16',
        'type': 'Danışmanlık', 'student_name': 'ADRIANA LIZZETTE CRUZ QUINTANILLA'
    }
    # Danışmanlık 3: Ersin Kaçmaz -> Salı 15:30 - 16:20 (Slot 7)
    teacher_schedule[9][1][7] = {
        'sube': '9', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 1, 'slot': 7, 'start': TIME_SLOTS[7][0], 'end': TIME_SLOTS[7][1],
        'room': room_9, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'ERSİN KAÇMAZ'
    }
    # Danışmanlık 4: Seçil Karaçalı -> Çarşamba 08:30 - 09:20 (Slot 0) - Seçil'in dersi 09:30'da başlıyor!
    teacher_schedule[9][2][0] = {
        'sube': '109', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 2, 'slot': 0, 'start': TIME_SLOTS[0][0], 'end': TIME_SLOTS[0][1],
        'room': room_9, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'SEÇİL KARAÇALI'
    }

    # 3. PROF. DR. RAMAZAN ÖZÇELİK (Şube 7) REVİZYONU
    # Amaç: Cuma = 0 saat, Pazartesi = 8 saat, Salı = 7 saat, Çarşamba = 7 saat, Perşembe = 6 saat.
    room_7 = TEACHERS[7]['room']
    
    # Lisans ve seçmeli dersleri koru:
    lisans_secmeli_7 = []
    for d in range(5):
        for sl in range(9):
            e = teacher_schedule[7][d][sl]
            if e and e['type'] in ('Lisans', 'Seçmeli'):
                lisans_secmeli_7.append((d, sl, e))
            teacher_schedule[7][d][sl] = None
            
    for d, sl, e in lisans_secmeli_7:
        teacher_schedule[7][d][sl] = e

    # Seminer Dersi: İNM 7006 (Çarşamba 08:30 - 10:20 - Slot 0, 1)
    for sl in [0, 1]:
        teacher_schedule[7][2][sl] = {
            'sube': '7', 'code': 'İNM 7006', 'name': 'Seminer II',
            'day': 2, 'slot': sl, 'start': TIME_SLOTS[sl][0], 'end': TIME_SLOTS[sl][1],
            'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '16',
            'type': 'Seminer', 'student_name': 'MAHAMAT WARDOUGOU LONY'
        }

    # Uzmanlık Alan Dersleri:
    # 1. Blok (DR - FBE 8901): Pazartesi 08:30 - 12:20 (Slot 0, 1, 2, 3)
    for sl in [0, 1, 2, 3]:
        teacher_schedule[7][0][sl] = {
            'sube': '7', 'code': 'FBE 8901', 'name': 'Uzmanlık Alan Dersi',
            'day': 0, 'slot': sl, 'start': TIME_SLOTS[sl][0], 'end': TIME_SLOTS[sl][1],
            'room': room_7, 'uyg': 0, 'ortak': 0, 'ckd': 0, 'prog': '16',
            'type': 'Uzmanlık', 'student_name': 'MAHAMAT WARDOUGOU LONY'
        }
    # 2. Blok (YL - FBE 6901): Çarşamba 13:30 - 17:20 (Slot 5, 6, 7, 8)
    yl_names_7 = "BUĞRA KÜÇÜK, HACI TIKNA, ERSİN KARAMAN, YUNUS ÇİFTÇİ, BEYZA AYTAÇ, DAMLA FİDANCI, GÜRSEL SEHA GÜLTEKİN"
    for sl in [5, 6, 7, 8]:
        teacher_schedule[7][2][sl] = {
            'sube': '7', 'code': 'FBE 6901', 'name': 'Uzmanlık Alan Dersi',
            'day': 2, 'slot': sl, 'start': TIME_SLOTS[sl][0], 'end': TIME_SLOTS[sl][1],
            'room': room_7, 'uyg': 0, 'ortak': 0, 'ckd': 0, 'prog': '15',
            'type': 'Uzmanlık', 'student_name': yl_names_7
        }

    # 10 Danışmanlık Dağılımı:
    # Pazartesi Öğleden Sonra (Pzt: 4 Uzmanlık + 1 İNM 403 + 3 Danışmanlık = 8 saat)
    teacher_schedule[7][0][5] = {
        'sube': '7', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 0, 'slot': 5, 'start': TIME_SLOTS[5][0], 'end': TIME_SLOTS[5][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'GÜRSEL SEHA GÜLTEKİN'
    }
    teacher_schedule[7][0][6] = {
        'sube': '107', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 0, 'slot': 6, 'start': TIME_SLOTS[6][0], 'end': TIME_SLOTS[6][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'HACI TIKNA'
    }
    teacher_schedule[7][0][7] = {
        'sube': '207', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 0, 'slot': 7, 'start': TIME_SLOTS[7][0], 'end': TIME_SLOTS[7][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'ERSİN KARAMAN'
    }

    # Salı (Salı: 1 Danışmanlık + 3 İNM 5059 + 1 İNM 403 + 2 Danışmanlık = 7 saat)
    teacher_schedule[7][1][0] = {
        'sube': '307', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 1, 'slot': 0, 'start': TIME_SLOTS[0][0], 'end': TIME_SLOTS[0][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'YUNUS ÇİFTÇİ'
    }
    teacher_schedule[7][1][5] = {
        'sube': '407', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 1, 'slot': 5, 'start': TIME_SLOTS[5][0], 'end': TIME_SLOTS[5][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'DAMLA FİDANCI'
    }
    teacher_schedule[7][1][6] = {
        'sube': '507', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 1, 'slot': 6, 'start': TIME_SLOTS[6][0], 'end': TIME_SLOTS[6][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'BEYZA AYTAÇ'
    }

    # Çarşamba (Çar: 2 Seminer + 1 Danışmanlık + 4 Uzmanlık = 7 saat)
    teacher_schedule[7][2][4] = {
        'sube': '7', 'code': 'FBE 7903', 'name': 'Danışmanlık',
        'day': 2, 'slot': 4, 'start': TIME_SLOTS[4][0], 'end': TIME_SLOTS[4][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '16',
        'type': 'Danışmanlık', 'student_name': 'SANAN GASIMOV'
    }

    # Perşembe (Per: 1 Danışmanlık + 1 Danışmanlık + 3 İNM 453 + 1 Danışmanlık = 6 saat)
    teacher_schedule[7][3][0] = {
        'sube': '7', 'code': 'FBE 8903', 'name': 'Danışmanlık',
        'day': 3, 'slot': 0, 'start': TIME_SLOTS[0][0], 'end': TIME_SLOTS[0][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '16',
        'type': 'Danışmanlık', 'student_name': 'MAHAMAT WARDOUGOU LONY'
    }
    teacher_schedule[7][3][4] = {
        'sube': '7', 'code': 'FBE 5903', 'name': 'Danışmanlık',
        'day': 3, 'slot': 4, 'start': TIME_SLOTS[4][0], 'end': TIME_SLOTS[4][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'VEYSEL AKIN'
    }
    teacher_schedule[7][3][8] = {
        'sube': '607', 'code': 'FBE 6903', 'name': 'Danışmanlık',
        'day': 3, 'slot': 8, 'start': TIME_SLOTS[8][0], 'end': TIME_SLOTS[8][1],
        'room': room_7, 'uyg': 1, 'ortak': 0, 'ckd': 0, 'prog': '15',
        'type': 'Danışmanlık', 'student_name': 'BUĞRA KÜÇÜK'
    }

    # Cuma Günü: 0 saat! (Tümü None kalır).

    # 4. EXCEL DOSYALARINI OLUŞTURMA
    
    # A. Bireysel Onay Dosyası: Prof. Dr. Ferhat ERDAL
    wb_ferhat = openpyxl.Workbook()
    ws_ferhat = wb_ferhat.active
    ws_ferhat.title = "Ferhat_ERDAL_Onay"
    populate_teacher_sheet(ws_ferhat, 9, teacher_schedule)
    ferhat_path = os.path.join(output_dir, "DersProgrami_Sube09_Prof_Dr_Ferhat_ERDAL_Onay.xlsx")
    wb_ferhat.save(ferhat_path)
    print(f"Oluşturuldu: {ferhat_path}")

    # B. Bireysel Onay Dosyası: Prof. Dr. Ramazan ÖZÇELİK
    wb_ramazan = openpyxl.Workbook()
    ws_ramazan = wb_ramazan.active
    ws_ramazan.title = "Ramazan_OZCELIK_Onay"
    populate_teacher_sheet(ws_ramazan, 7, teacher_schedule)
    ramazan_path = os.path.join(output_dir, "DersProgrami_Sube07_Prof_Dr_Ramazan_OZCELIK_Onay.xlsx")
    wb_ramazan.save(ramazan_path)
    print(f"Oluşturuldu: {ramazan_path}")

    # C. Genel Bölüm Özeti ve 15 Hoca Revize Master Dosyası:
    master_path = os.path.join(output_dir, "DersProgrami_Kisisel_OnayRevize.xlsx")
    
    wb_master = openpyxl.Workbook()
    ws_summary = wb_master.active
    ws_summary.title = "Bölüm Genel Özeti"
    ws_summary.views.sheetView[0].showGridLines = True

    # Başlık
    ws_summary.merge_cells("A1:K1")
    ws_summary["A1"] = "AKDENİZ ÜNİVERSİTESİ İNŞAAT MÜHENDİSLİĞİ BÖLÜMÜ"
    ws_summary["A1"].font = Font(name='Segoe UI', size=16, bold=True, color='1E293B')
    ws_summary["A1"].alignment = Alignment(horizontal='center', vertical='center')

    ws_summary.merge_cells("A2:K2")
    ws_summary["A2"] = "Öğretim Üyeleri Haftalık Ders Yükü ve Görevlendirme Özeti (Revize Onay Taslağı)"
    ws_summary["A2"].font = Font(name='Segoe UI', size=11, color='64748B')
    ws_summary["A2"].alignment = Alignment(horizontal='center', vertical='center')

    headers_summary = [
        "Şube", "Öğretim Üyesi", "Oda / Derslik", "Lisans (Saat)", "YL/DR Seçmeli",
        "Uzmanlık Alan", "Danışmanlık", "Seminer", "Bölüm İçi Toplam",
        "Tezsiz YL (İ.Ö.)*", "Genel Toplam"
    ]
    ws_summary.append([])
    ws_summary.append(headers_summary)
    
    header_fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx in range(1, 12):
        cell = ws_summary.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws_summary.row_dimensions[4].height = 28

    current_r = 5
    for t_id in sorted(TEACHERS.keys()):
        t_info = TEACHERS[t_id]
        l_cnt = sum(1 for d in range(5) for sl in range(9) if teacher_schedule[t_id][d][sl] and teacher_schedule[t_id][d][sl]['type'] == 'Lisans')
        s_cnt = sum(1 for d in range(5) for sl in range(9) if teacher_schedule[t_id][d][sl] and teacher_schedule[t_id][d][sl]['type'] == 'Seçmeli')
        u_cnt = sum(1 for d in range(5) for sl in range(9) if teacher_schedule[t_id][d][sl] and teacher_schedule[t_id][d][sl]['type'] == 'Uzmanlık')
        d_cnt = sum(1 for d in range(5) for sl in range(9) if teacher_schedule[t_id][d][sl] and teacher_schedule[t_id][d][sl]['type'] == 'Danışmanlık')
        sem_cnt = sum(1 for d in range(5) for sl in range(9) if teacher_schedule[t_id][d][sl] and teacher_schedule[t_id][d][sl]['type'] == 'Seminer')
        bolum_tot = l_cnt + s_cnt + u_cnt + d_cnt + sem_cnt
        eve_tot = sum(c['hours'] for c in SECOND_EDUCATION_COURSES.get(t_id, []))
        genel_tot = bolum_tot + eve_tot

        row_vals = [t_id, t_info['title'], t_info['room'], l_cnt, s_cnt, u_cnt, d_cnt, sem_cnt, bolum_tot, eve_tot, genel_tot]
        ws_summary.append(row_vals)
        
        row_fill = PatternFill(start_color='F8FAFC' if current_r % 2 == 0 else 'FFFFFF', end_color='F8FAFC' if current_r % 2 == 0 else 'FFFFFF', fill_type='solid')
        for col_idx in range(1, 12):
            cell = ws_summary.cell(row=current_r, column=col_idx)
            cell.fill = row_fill
            cell.font = Font(name='Segoe UI', size=10, bold=(col_idx in [2, 9, 11]))
            cell.alignment = Alignment(horizontal='left' if col_idx == 2 else 'center', vertical='center')
            cell.border = thin_border
        ws_summary.row_dimensions[current_r].height = 22
        current_r += 1

    # Toplam Satırı
    ws_summary.append([
        "", "BÖLÜM TOPLAMI", "",
        f"=SUM(D5:D{current_r-1})", f"=SUM(E5:E{current_r-1})", f"=SUM(F5:F{current_r-1})",
        f"=SUM(G5:G{current_r-1})", f"=SUM(H5:H{current_r-1})", f"=SUM(I5:I{current_r-1})",
        f"=SUM(J5:J{current_r-1})", f"=SUM(K5:K{current_r-1})"
    ])
    for col_idx in range(1, 12):
        cell = ws_summary.cell(row=current_r, column=col_idx)
        cell.fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
        cell.font = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
        cell.alignment = Alignment(horizontal='left' if col_idx == 2 else 'center', vertical='center')
        cell.border = thin_border
    ws_summary.row_dimensions[current_r].height = 24

    ws_summary.column_dimensions['A'].width = 8
    ws_summary.column_dimensions['B'].width = 34
    ws_summary.column_dimensions['C'].width = 16
    ws_summary.column_dimensions['D'].width = 15
    ws_summary.column_dimensions['E'].width = 16
    ws_summary.column_dimensions['F'].width = 16
    ws_summary.column_dimensions['G'].width = 15
    ws_summary.column_dimensions['H'].width = 14
    ws_summary.column_dimensions['I'].width = 18
    ws_summary.column_dimensions['J'].width = 18
    ws_summary.column_dimensions['K'].width = 16

    # 15 Hoca Sayfaları
    for t_id in sorted(TEACHERS.keys()):
        t_info = TEACHERS[t_id]
        clean_title = t_info['title'].replace('Prof. Dr. ', 'Prof.Dr.').replace('Doç. Dr. ', 'Doç.Dr.').replace('Dr. Öğr. Üyesi ', 'Dr.Öğr.Üyesi ')
        sheet_name = clean_title[:31]
        ws = wb_master.create_sheet(title=sheet_name)
        populate_teacher_sheet(ws, t_id, teacher_schedule)

    wb_master.save(master_path)
    print(f"Oluşturuldu: {master_path}")

    # Doğrulama Raporu
    print("\n--- GÜNLÜK DAĞILIM KONTROLÜ ---")
    for t_id in [7, 9]:
        daily = [sum(1 for sl in range(9) if teacher_schedule[t_id][d][sl] is not None) for d in range(5)]
        print(f"Şube {t_id} ({TEACHERS[t_id]['title']}): Pzt:{daily[0]}, Sal:{daily[1]}, Çar:{daily[2]}, Per:{daily[3]}, Cum:{daily[4]} | Toplam: {sum(daily)} | Max: {max(daily)}")

if __name__ == '__main__':
    generate_revised_onay_schedules()
